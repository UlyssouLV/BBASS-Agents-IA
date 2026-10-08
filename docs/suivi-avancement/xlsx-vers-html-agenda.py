#!/usr/bin/env python3
"""Lit un planning Excel (feuille Planning) et écrit sur la sortie standard un HTML table prêt à capturer."""
from __future__ import annotations

import datetime as dt
import html
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl manquant : pip install openpyxl"
    ) from exc


def fmt_valeur(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, dt.datetime):
        return valeur.strftime("%d/%m/%Y")
    if isinstance(valeur, dt.time):
        return f"{valeur.hour}:{valeur.minute:02d}"
    if isinstance(valeur, dt.date):
        return valeur.strftime("%d/%m/%Y")
    texte = str(valeur).replace("\xa0", " ").strip()
    return texte


def couleur_theme(theme: int | None, tint: float | None) -> str | None:
    """Approx. des thèmes Office courants → hex (suffisant pour le PDF)."""
    base = {
        0: (255, 255, 255),
        1: (0, 0, 0),
        2: (238, 236, 225),
        3: (31, 73, 125),
        4: (79, 129, 189),
        5: (192, 80, 77),
        6: (155, 187, 89),
        7: (128, 100, 162),
        8: (75, 172, 198),
        9: (247, 150, 70),
    }
    if theme is None or theme not in base:
        return None
    r, g, b = base[theme]
    t = float(tint or 0.0)
    if t > 0:
        r = int(r + (255 - r) * t)
        g = int(g + (255 - g) * t)
        b = int(b + (255 - b) * t)
    elif t < 0:
        r = int(r * (1 + t))
        g = int(g * (1 + t))
        b = int(b * (1 + t))
    return f"#{r:02x}{g:02x}{b:02x}"


def fill_css(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type not in ("solid", "patternFill"):
        return ""
    fg = fill.fgColor
    if fg is None:
        return ""
    if getattr(fg, "type", None) == "rgb" and fg.rgb and fg.rgb != "00000000":
        rgb = fg.rgb[-6:]
        return f"background:#{rgb};"
    if getattr(fg, "type", None) == "theme":
        hex_c = couleur_theme(fg.theme, getattr(fg, "tint", None))
        if hex_c:
            return f"background:{hex_c};"
    return ""


def font_css(cell) -> str:
    font = cell.font
    parts = []
    if font.bold:
        parts.append("font-weight:700;")
    if font.size:
        # Excel pt → CSS px approx. pour capture dense
        size = float(font.size)
        if size >= 48:
            parts.append("font-size:28px;")
        elif size >= 20:
            parts.append("font-size:16px;")
        elif size >= 14:
            parts.append("font-size:11px;")
        else:
            parts.append("font-size:10px;")
    color = font.color
    if color is not None:
        if getattr(color, "type", None) == "rgb" and color.rgb:
            parts.append(f"color:#{color.rgb[-6:]};")
        elif getattr(color, "type", None) == "theme":
            hex_c = couleur_theme(color.theme, getattr(color, "tint", None))
            if hex_c:
                parts.append(f"color:{hex_c};")
    return "".join(parts)


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage : python xlsx-vers-html-agenda.py <source.xlsx> > sortie.html"
        )
    source = Path(sys.argv[1])
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb["Planning"] if "Planning" in wb.sheetnames else wb.active

    # Zone utile : colonnes B–G
    min_col, max_col = 2, 7
    n_cols = max_col - min_col + 1

    merges: dict[tuple[int, int], tuple[int, int]] = {}
    skip: set[tuple[int, int]] = set()
    for plage in ws.merged_cells.ranges:
        min_r, min_c, max_r, max_c = plage.min_row, plage.min_col, plage.max_row, plage.max_col
        # Hors zone utile : ignorer
        if max_c < min_col or min_c > max_col or max_r < 2:
            continue
        min_c = max(min_c, min_col)
        max_c = min(max_c, max_col)
        rowspan = max_r - min_r + 1
        colspan = max_c - min_c + 1
        # Fusion pleine largeur : pas de rowspan HTML (sinon la ligne suivante
        # est poussée en colonnes 7+ → débordement à droite).
        if min_c == min_col and colspan == n_cols and rowspan > 1:
            rowspan = 1
            max_r = min_r
        merges[(min_r, min_c)] = (rowspan, colspan)
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                if (r, c) != (min_r, min_c):
                    skip.add((r, c))

    max_row = 2
    for r in range(2, (ws.max_row or 2) + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(min_col, max_col + 1)):
            max_row = r

    def row_has_content(r: int) -> bool:
        for c in range(min_col, max_col + 1):
            if (r, c) in skip:
                continue
            if (r, c) in merges:
                return True
            if ws.cell(r, c).value not in (None, ""):
                return True
        return False

    # Ajuste les rowspan si on saute des lignes vides au milieu d’une fusion.
    skipped_rows: set[int] = set()
    for r in range(2, max_row + 1):
        if not row_has_content(r):
            skipped_rows.add(r)

    for (r0, c0), (rs, cs) in list(merges.items()):
        if rs <= 1:
            continue
        kept = sum(1 for r in range(r0, r0 + rs) if r not in skipped_rows)
        if kept < 1:
            kept = 1
        merges[(r0, c0)] = (kept, cs)

    lignes = []
    for r in range(2, max_row + 1):
        if r in skipped_rows:
            continue

        # Ligne « SEMAINE : date » éparpillée → une seule cellule pleine largeur
        valeurs_utiles = [
            (c, ws.cell(r, c).value)
            for c in range(min_col, max_col + 1)
            if (r, c) not in skip and ws.cell(r, c).value not in (None, "")
        ]
        if (
            r not in (6,)
            and (r, min_col) not in merges
            and len(valeurs_utiles) >= 1
            and all(
                isinstance(v, (str, dt.datetime, dt.date)) or v is None
                for _, v in (
                    (c, ws.cell(r, c).value) for c in range(min_col, max_col + 1)
                )
            )
            and not any(isinstance(ws.cell(r, c).value, dt.time) for c in range(min_col, max_col + 1))
            and all(c > min_col for c, _ in valeurs_utiles)  # pas une ligne horaire
        ):
            parts = [fmt_valeur(v) for _, v in valeurs_utiles]
            texte = html.escape(" ".join(parts))
            style = (
                "border:1px solid #94a3b8;padding:6px 8px;vertical-align:middle;"
                "text-align:center;font-weight:650;box-sizing:border-box;"
            )
            lignes.append(
                f'<tr><td colspan="{n_cols}" style="{style}">{texte}</td></tr>'
            )
            continue

        cells_html = []
        for c in range(min_col, max_col + 1):
            if (r, c) in skip:
                continue
            cell = ws.cell(r, c)
            rowspan, colspan = merges.get((r, c), (1, 1))
            texte = html.escape(fmt_valeur(cell.value)).replace("\n", "<br />")
            style = (
                "border:1px solid #94a3b8;padding:4px 6px;vertical-align:middle;"
                "text-align:center;word-break:break-word;box-sizing:border-box;"
                + fill_css(cell)
                + font_css(cell)
            )
            # En-têtes jour / heure : forcer contraste lisible
            if r == 6:
                style += "background:#1e3a5f;color:#fff;font-weight:700;"
            attrs = f' style="{style}"'
            if rowspan > 1:
                attrs += f' rowspan="{rowspan}"'
            if colspan > 1:
                attrs += f' colspan="{colspan}"'
            tag = "th" if r == 6 else "td"
            cells_html.append(f"<{tag}{attrs}>{texte}</{tag}>")
        if cells_html:
            lignes.append("<tr>" + "".join(cells_html) + "</tr>")

    doc = f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8" />
<title>Emploi du temps</title>
<style>
  html, body {{ margin: 0; padding: 0; background: #fff; }}
  table {{
    border-collapse: collapse;
    border-spacing: 0;
    table-layout: fixed;
    width: 1170px;
    max-width: 1170px;
    font-family: "Segoe UI", Calibri, sans-serif;
    font-size: 10.5px;
    line-height: 1.25;
    color: #1c1917;
  }}
  col.heure {{ width: 70px; }}
  col.jour {{ width: 220px; }}
  td, th {{ box-sizing: border-box; }}
</style>
</head>
<body>
<table id="agenda">
  <colgroup>
    <col class="heure" />
    <col class="jour" /><col class="jour" /><col class="jour" />
    <col class="jour" /><col class="jour" />
  </colgroup>
  <tbody>
    {"".join(lignes)}
  </tbody>
</table>
</body>
</html>
"""
    # Sortie standard plutôt qu'un chemin en argument : l'appelant choisit où écrire.
    sys.stdout.buffer.write(doc.encode("utf-8"))


if __name__ == "__main__":
    main()
