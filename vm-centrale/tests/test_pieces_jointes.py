from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from vm_centrale.models import PieceJointe

_TYPE_MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TYPE_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fichier_docx(*paragraphes: str) -> bytes:
    document = Document()
    for paragraphe in paragraphes:
        document.add_paragraph(paragraphe)
    tampon = BytesIO()
    document.save(tampon)
    return tampon.getvalue()


def _fichier_xlsx(lignes: list[list[object]]) -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    assert feuille is not None
    for ligne in lignes:
        feuille.append(ligne)
    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def _autorisation(jeton: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {jeton}"}


def _creer_conversation(client, mistral_client_factice, jeton: str, message: str = "Bonjour") -> int:
    mistral_client_factice.repondre("Réponse assistant", "Titre")
    reponse = client.post(
        "/conversations", json={"message": message}, headers=_autorisation(jeton)
    )
    return reponse.json()["conversation"]["id"]


def _jeton_admin(client, seed_compte, identifiant: str = "a.martin") -> str:
    seed_compte(
        identifiant, "correcthorsebatterystaple", agence="Castries", poles=["Foncier"],
        prenom="Alice", nom="Martin", est_admin=True,
    )
    reponse = client.post(
        "/auth", json={"identifiant": identifiant, "mot_de_passe": "correcthorsebatterystaple"}
    )
    return reponse.json()["jeton"]


@pytest.fixture(autouse=True)
def _repertoire_pieces_jointes(tmp_path, monkeypatch):
    # Isole les tests du répertoire par défaut (VM_CENTRALE_PIECES_JOINTES_DIR,
    # ./pieces_jointes) : chaque test écrit dans un répertoire temporaire
    # propre, jamais dans le dépôt.
    monkeypatch.setattr("vm_centrale.routers.conversations.PIECES_JOINTES_DIR", str(tmp_path))
    return tmp_path


def test_upload_pdf_valide_renvoie_201_stocke_le_fichier_et_extrait_le_contenu(
    client, mistral_client_factice, jeton_valide, db_session, _repertoire_pieces_jointes
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    mistral_client_factice.repondre_ocr("Texte extrait du PDF")

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("document.pdf", b"%PDF-1.4 contenu factice", "application/pdf")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 201
    corps = reponse.json()
    assert corps["piece_jointe"]["nom_fichier"] == "document.pdf"
    assert corps["piece_jointe"]["type_mime"] == "application/pdf"
    assert corps["echec_analyse"] is False

    piece_jointe_id = corps["piece_jointe"]["id"]
    piece_jointe = db_session.get(PieceJointe, piece_jointe_id)
    assert piece_jointe.contenu_extrait == "Texte extrait du PDF"
    assert piece_jointe.conversation_id == conversation_id
    assert piece_jointe.message_id is None

    chemin_attendu = Path(_repertoire_pieces_jointes) / "j.dupont" / str(conversation_id) / f"{piece_jointe_id}-document.pdf"
    assert chemin_attendu.exists()
    assert chemin_attendu.read_bytes() == b"%PDF-1.4 contenu factice"

    assert mistral_client_factice.appels_ocr == [(b"%PDF-1.4 contenu factice", "application/pdf")]


def test_upload_docx_valide_renvoie_201_et_extrait_le_texte_sans_appeler_mistral(
    client, mistral_client_factice, jeton_valide, db_session, _repertoire_pieces_jointes
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    contenu_docx = _fichier_docx("Première ligne du document", "Seconde ligne")
    appels_chat_avant = len(mistral_client_factice.messages_recus)

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("document.docx", contenu_docx, _TYPE_MIME_DOCX)},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 201
    corps = reponse.json()
    assert corps["piece_jointe"]["nom_fichier"] == "document.docx"
    assert corps["piece_jointe"]["type_mime"] == _TYPE_MIME_DOCX
    assert corps["echec_analyse"] is False

    piece_jointe = db_session.get(PieceJointe, corps["piece_jointe"]["id"])
    assert piece_jointe.contenu_extrait == "Première ligne du document\nSeconde ligne"

    chemin_attendu = (
        Path(_repertoire_pieces_jointes)
        / "j.dupont"
        / str(conversation_id)
        / f"{piece_jointe.id}-document.docx"
    )
    assert chemin_attendu.read_bytes() == contenu_docx

    # Aucun format sans alternative locale ici : ni .chat() ni .ocr() ne
    # doivent être sollicités pour du Word (spec 1.1.2). Compare un avant/
    # après plutôt qu'une liste vide : _creer_conversation a déjà appelé
    # .chat() deux fois (réponse + titrage) avant l'upload.
    assert mistral_client_factice.appels_ocr == []
    assert len(mistral_client_factice.messages_recus) == appels_chat_avant


def test_upload_xlsx_valide_renvoie_201_et_extrait_les_valeurs_de_cellules_sans_appeler_mistral(
    client, mistral_client_factice, jeton_valide, db_session, _repertoire_pieces_jointes
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    contenu_xlsx = _fichier_xlsx([["Nom", "Montant"], ["Dupont", 42]])
    appels_chat_avant = len(mistral_client_factice.messages_recus)

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("tableau.xlsx", contenu_xlsx, _TYPE_MIME_XLSX)},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 201
    corps = reponse.json()
    assert corps["piece_jointe"]["nom_fichier"] == "tableau.xlsx"
    assert corps["piece_jointe"]["type_mime"] == _TYPE_MIME_XLSX
    assert corps["echec_analyse"] is False

    piece_jointe = db_session.get(PieceJointe, corps["piece_jointe"]["id"])
    assert piece_jointe.contenu_extrait == "Nom\tMontant\nDupont\t42"

    assert mistral_client_factice.appels_ocr == []
    assert len(mistral_client_factice.messages_recus) == appels_chat_avant


def test_upload_image_valide_renvoie_201_et_extrait_le_contenu_via_lappel_vision(
    client, mistral_client_factice, jeton_valide, db_session, _repertoire_pieces_jointes
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    mistral_client_factice.repondre_vision("Photocopie d'un plan de masse, échelle 1/200")

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("plan.png", b"contenu factice png", "image/png")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 201
    corps = reponse.json()
    assert corps["piece_jointe"]["nom_fichier"] == "plan.png"
    assert corps["piece_jointe"]["type_mime"] == "image/png"
    assert corps["echec_analyse"] is False

    piece_jointe = db_session.get(PieceJointe, corps["piece_jointe"]["id"])
    assert piece_jointe.contenu_extrait == "Photocopie d'un plan de masse, échelle 1/200"
    assert piece_jointe.echec_analyse is False

    # L'appel vision passe par .chat() avec response_format (pas .ocr(),
    # jamais utilisé pour une image) : aucun appel OCR ne doit être fait.
    assert mistral_client_factice.appels_ocr == []


def test_upload_image_dont_lanalyse_echoue_renvoie_201_avec_echec_analyse_a_true(
    client, mistral_client_factice, jeton_valide, db_session
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    mistral_client_factice.repondre_vision(
        "Image illisible : trop floue pour en extraire quoi que ce soit d'exploitable",
        echec_analyse=True,
    )

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("photo.jpg", b"contenu factice jpg", "image/jpeg")},
        headers=_autorisation(jeton_valide),
    )

    # Un échec d'analyse (rien d'exploitable dans l'image) reste un succès
    # HTTP : jamais une erreur 4xx/5xx pour ce cas-là (spec 1.1.2).
    assert reponse.status_code == 201
    corps = reponse.json()
    assert corps["echec_analyse"] is True

    piece_jointe = db_session.get(PieceJointe, corps["piece_jointe"]["id"])
    assert piece_jointe.echec_analyse is True
    assert piece_jointe.contenu_extrait == (
        "Image illisible : trop floue pour en extraire quoi que ce soit d'exploitable"
    )


def test_upload_dune_conversation_dun_autre_compte_renvoie_404(
    client, mistral_client_factice, jeton_valide, jeton_store
):
    jeton_autre_compte = jeton_store.emettre("n.durand")
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_autre_compte)

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("document.pdf", b"contenu", "application/pdf")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 404


def test_upload_dune_conversation_dun_autre_compte_renvoie_404_meme_avec_un_jeton_admin(
    client, mistral_client_factice, jeton_valide, seed_compte
):
    jeton_admin = _jeton_admin(client, seed_compte)
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("document.pdf", b"contenu", "application/pdf")},
        headers=_autorisation(jeton_admin),
    )

    assert reponse.status_code == 404


def test_upload_dune_conversation_inexistante_renvoie_404(client, jeton_valide):
    reponse = client.post(
        "/conversations/999/pieces-jointes",
        files={"fichier": ("document.pdf", b"contenu", "application/pdf")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 404


def test_upload_sans_jeton_est_refuse(client):
    reponse = client.post(
        "/conversations/1/pieces-jointes",
        files={"fichier": ("document.pdf", b"contenu", "application/pdf")},
    )

    assert reponse.status_code == 401


def test_upload_fichier_trop_volumineux_renvoie_400(
    client, mistral_client_factice, jeton_valide
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    contenu_trop_gros = b"x" * (20 * 1024 * 1024 + 1)

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("document.pdf", contenu_trop_gros, "application/pdf")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 400
    assert mistral_client_factice.appels_ocr == []


def test_upload_type_non_supporte_renvoie_400(client, mistral_client_factice, jeton_valide):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("notes.txt", b"du texte brut", "text/plain")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 400
    assert mistral_client_factice.appels_ocr == []


def test_upload_echec_appel_ocr_renvoie_502_et_ne_persiste_aucune_piece_jointe(
    client, mistral_client_factice, jeton_valide, db_session
):
    conversation_id = _creer_conversation(client, mistral_client_factice, jeton_valide)
    mistral_client_factice.echouer_ocr(RuntimeError("service Mistral indisponible"))

    reponse = client.post(
        f"/conversations/{conversation_id}/pieces-jointes",
        files={"fichier": ("document.pdf", b"contenu", "application/pdf")},
        headers=_autorisation(jeton_valide),
    )

    assert reponse.status_code == 502
    assert db_session.query(PieceJointe).count() == 0
