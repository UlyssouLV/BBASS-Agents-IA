from io import BytesIO

from openpyxl import load_workbook

from vm_centrale.analyse_pieces_jointes.resultat import ResultatAnalyse
from vm_centrale.mistral_client import MistralClient


def analyser(fichier: bytes, _type_mime: str, _client: MistralClient) -> ResultatAnalyse:
    # Extraction locale (openpyxl), même logique que word.py : un tableur est
    # une donnée déjà structurée, restituée en texte/tableau, jamais transmis
    # à Mistral (spec 1.1.2). `_type_mime`/`_client` non utilisés ici, gardés
    # (préfixés `_`) pour la signature uniforme des sous-modules appelés de
    # façon polymorphe par analyse_pieces_jointes.analyser (voir __init__.py).
    classeur = load_workbook(BytesIO(fichier), data_only=True, read_only=True)
    lignes = []
    for feuille in classeur.worksheets:
        for ligne in feuille.iter_rows(values_only=True):
            valeurs = [str(valeur) for valeur in ligne if valeur is not None]
            if valeurs:
                lignes.append("\t".join(valeurs))
    contenu_extrait = "\n".join(lignes)
    return ResultatAnalyse(
        contenu_extrait=contenu_extrait, echec_analyse=not contenu_extrait.strip()
    )
